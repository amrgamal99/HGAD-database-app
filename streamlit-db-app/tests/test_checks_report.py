from pathlib import Path
from io import BytesIO
import sys
from types import SimpleNamespace

from openpyxl import load_workbook
import pandas as pd

PROJECT_SRC = Path(__file__).resolve().parents[1] / "src"
sys.path.insert(0, str(PROJECT_SRC))

from db.connection import fetch_checks_report_data
from reports.checks import (
    CHECK_COLUMNS,
    PURPOSE_FILTER_OPTIONS,
    _factory_cards_html,
    _factory_check_summary,
    _filter_checks_by_purpose,
    _prepare_display_dataframe,
    _to_excel_bytes,
)


def test_check_details_have_requested_columns_order_and_clean_dates():
    source = pd.DataFrame(
        {
            "factoryname": ["بدر"],
            "companyname": ["شركة مثال"],
            "اسم المشروع": ["عقد تجريبي"],
            "رقم الشيك": ["123"],
            "تاريخ الشيك": ["2026-10-03 00:00:00"],
            "البنك": ["بنك مثال"],
            "قيمة الشيك": [250],
            "رابط الشيك": ["https://example.test/check"],
            "الغرض من إصدار الشيك": ["دفعة مقدمة"],
            "تاريخ استلام شيك": ["2026-10-04"],
            "السنة_المالية_المحاسبية": [2026],
        }
    )

    display_df = _prepare_display_dataframe(source)

    assert display_df.columns.tolist() == list(CHECK_COLUMNS)
    assert display_df.loc[0, "اسم المصنع"] == "بدر"
    assert display_df.loc[0, "اسم العقد"] == "عقد تجريبي"
    assert display_df.loc[0, "رابط نسخة الشيك"] == "https://example.test/check"
    assert display_df.loc[0, "تاريخ الشيك"] == "2026-10-03"


def test_factory_summary_totals_and_normalizes_purpose_variants():
    source = pd.DataFrame(
        {
            "اسم المصنع": ["بدر", "بدر", "بدر", "التجمع"],
            "قيمة الشيك": [100, 200, 300, 50],
            "الغرض من إصدار الشيك": [
                "دفعه مقدمه",
                "رد تأمينات اجتماعية",
                "سداد مستخلص",
                "سداد مستخلص",
            ],
        }
    )

    summary = _factory_check_summary(source).set_index("اسم المصنع")

    assert summary.loc["بدر", "الإجمالي"] == 600
    assert summary.loc["بدر", "دفعة مقدمة"] == 100
    assert summary.loc["بدر", "رد تامينات اجتماعيه"] == 200
    assert summary.loc["بدر", "سداد مستخلص"] == 300


def test_check_purpose_filter_matches_the_three_supported_values():
    assert PURPOSE_FILTER_OPTIONS == (
        "رد تامينات اجتماعيه",
        "سداد مستخلص",
        "دفعه مقدمه",
    )

    source = pd.DataFrame(
        {
            "رقم الشيك": ["1", "2", "3"],
            "الغرض من إصدار الشيك": [
                "دفعة مقدمة",
                "رد تأمينات اجتماعية",
                "سداد مستخلص",
            ],
        }
    )

    filtered = _filter_checks_by_purpose(source, "دفعه مقدمه")

    assert filtered["رقم الشيك"].tolist() == ["1"]


def test_check_number_prefix_is_sent_as_a_literal_ilike_prefix():
    queries = {}

    class Query:
        def __init__(self, table_name):
            self.table_name = table_name
            self.filters = []

        def select(self, *_columns):
            return self

        def ilike(self, column, pattern):
            self.filters.append(("ilike", column, pattern))
            return self

        def execute(self):
            data = [{"رقم الشيك": "12%34"}] if self.table_name == "checks" else []
            return SimpleNamespace(data=data)

    class Client:
        def table(self, table_name):
            query = Query(table_name)
            queries[table_name] = query
            return query

    result = fetch_checks_report_data(
        Client(),
        date_field="تاريخ الشيك",
        check_number_prefix="12%",
    )

    assert result["رقم الشيك"].tolist() == ["12%34"]
    assert queries["checks"].filters == [("ilike", "رقم الشيك", r"12\%%")]


def test_factory_cards_render_full_amounts_without_ellipsis():
    summary = pd.DataFrame(
        [
            {
                "اسم المصنع": "بدر",
                "الإجمالي": 262520893.79,
                "دفعة مقدمة": 291.90,
                "رد تامينات اجتماعيه": 659.96,
                "سداد مستخلص": 941.93,
            }
        ]
    )

    card_html = _factory_cards_html(summary)

    assert "262,520,893.79" in card_html
    assert "291.90" in card_html
    assert "659.96" in card_html
    assert "941.93" in card_html
    assert "..." not in card_html


def test_checks_excel_is_formatted_as_rtl_table_with_clickable_anchor():
    source = pd.DataFrame(
        {
            "اسم المصنع": ["بدر"],
            "اسم الشركة": ["شركة مثال"],
            "اسم العقد": ["عقد تجريبي"],
            "رقم الشيك": ["123"],
            "تاريخ الشيك": ["2026-10-03"],
            "البنك": ["بنك مثال"],
            "قيمة الشيك": [1234.5],
            "رابط نسخة الشيك": ["https://example.test/check"],
            "الغرض من إصدار الشيك": ["دفعة مقدمة"],
            "تاريخ استلام شيك": ["2026-10-04"],
            "السنة_المالية_المحاسبية": [2026],
        }
    )

    workbook = load_workbook(BytesIO(_to_excel_bytes(source)))
    worksheet = workbook["الشيكات"]

    assert worksheet["A1"].value == "تقرير الشيكات"
    assert worksheet.sheet_view.rightToLeft is True
    assert worksheet.freeze_panes == "A7"
    assert worksheet.tables["ChecksTable"].ref == "A6:K7"
    assert worksheet.tables["ChecksTable"].tableStyleInfo.name == "TableStyleMedium2"
    assert worksheet["H7"].value == "فتح الشيك"
    assert worksheet["H7"].hyperlink.target == "https://example.test/check"
